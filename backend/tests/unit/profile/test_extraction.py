from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document

from app.core.errors import UnprocessableError
from app.modules.profile.extraction import FileKind, detect_kind, extract

FIXTURES = Path(__file__).parents[2] / "fixtures" / "cv"


def test_pdf_text_layer() -> None:
    data = (FIXTURES / "sample_cv.pdf").read_bytes()
    assert detect_kind(data, "cv.pdf") == FileKind.PDF
    extracted = extract(data, FileKind.PDF)
    assert "Nur Aina Rahman" in extracted.text
    assert "Selat Pay" in extracted.text
    assert not extracted.needs_vision


def test_image_goes_to_vision() -> None:
    data = (FIXTURES / "sample_cv.png").read_bytes()
    assert detect_kind(data, "scan.png") == FileKind.IMAGE
    extracted = extract(data, FileKind.IMAGE)
    assert extracted.needs_vision
    assert extracted.images[0].startswith(b"\x89PNG")


def test_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Aina Rahman")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "FastAPI"
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    assert detect_kind(data, "cv.docx") == FileKind.DOCX
    text = extract(data, FileKind.DOCX).text
    assert "Aina Rahman" in text
    assert "Python | FastAPI" in text


def test_plain_text() -> None:
    data = b"Aina Rahman\n\n\n\nEngineer   at   Selat"
    assert detect_kind(data, "cv.txt") == FileKind.TEXT
    assert extract(data, FileKind.TEXT).text == "Aina Rahman\n\nEngineer at Selat"


@pytest.mark.parametrize(
    ("data", "name"), [(b"GIF89a....", "cv.gif"), (b"\x00\x01binary", "cv.exe")]
)
def test_unsupported_files_are_refused(data: bytes, name: str) -> None:
    with pytest.raises(UnprocessableError):
        detect_kind(data, name)


def test_broken_pdf_is_refused() -> None:
    with pytest.raises(UnprocessableError):
        extract(b"%PDF-1.7 not really a pdf", FileKind.PDF)
