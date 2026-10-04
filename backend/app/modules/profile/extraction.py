"""Get plain text out of an uploaded CV.

PDF (text layer via pdfium), Word (.docx), plain text, and images. Scanned PDFs
and images have no text layer: they are rendered to images for the vision model
(see `service.read_with_vision`).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from enum import StrEnum

import filetype
import pypdfium2 as pdfium
from docx import Document
from PIL import Image

from app.core.errors import UnprocessableError

MAX_PAGES = 8
VISION_PAGES = 4
MIN_TEXT_CHARS = 200  # less than this from a PDF means it is scanned


class FileKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    TEXT = "txt"
    IMAGE = "image"


_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_IMAGE_MIMES = {"image/png": "png", "image/jpeg": "jpg", "image/webp": "webp"}


@dataclass
class Extracted:
    kind: FileKind
    text: str
    pages: int = 0
    images: list[bytes] = field(default_factory=list)  # PNGs for the vision model

    @property
    def needs_vision(self) -> bool:
        return bool(self.images) and len(self.text) < MIN_TEXT_CHARS


def detect_kind(data: bytes, filename: str) -> FileKind:
    """Trust the bytes, not the file name."""
    guess = filetype.guess(data)
    mime = guess.mime if guess else None
    if mime == "application/pdf":
        return FileKind.PDF
    if mime in _IMAGE_MIMES:
        return FileKind.IMAGE
    if mime in {_DOCX_MIME, "application/zip"} and filename.lower().endswith(".docx"):
        return FileKind.DOCX
    if mime is None and filename.lower().endswith((".txt", ".md")):
        try:
            data.decode("utf-8")
        except UnicodeDecodeError as error:
            raise UnprocessableError("That text file isn't readable.", code="bad_file") from error
        return FileKind.TEXT
    raise UnprocessableError(
        "Upload a PDF, Word (.docx), image (PNG, JPG) or text file.", code="unsupported_file"
    )


def extract(data: bytes, kind: FileKind) -> Extracted:
    match kind:
        case FileKind.PDF:
            return _from_pdf(data)
        case FileKind.DOCX:
            return Extracted(kind, _tidy(_from_docx(data)))
        case FileKind.TEXT:
            return Extracted(kind, _tidy(data.decode("utf-8", errors="replace")))
        case FileKind.IMAGE:
            return Extracted(kind, "", pages=1, images=[_normalise_image(data)])


def _from_pdf(data: bytes) -> Extracted:
    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as error:
        raise UnprocessableError(
            "That PDF couldn't be opened. Is it password-protected?", code="bad_file"
        ) from error
    try:
        pages = len(pdf)
        texts = []
        for index in range(min(pages, MAX_PAGES)):
            page = pdf[index]
            textpage = page.get_textpage()
            texts.append(textpage.get_text_bounded())
            textpage.close()
            page.close()
        text = _tidy("\n\n".join(texts))
        images: list[bytes] = []
        if len(text) < MIN_TEXT_CHARS:
            for index in range(min(pages, VISION_PAGES)):
                page = pdf[index]
                bitmap = page.render(scale=2.0)
                images.append(_png(bitmap.to_pil()))
                page.close()
        return Extracted(FileKind.PDF, text, pages=pages, images=images)
    finally:
        pdf.close()


def _from_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as error:  # python-docx raises several unrelated types
        raise UnprocessableError("That Word file couldn't be opened.", code="bad_file") from error
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


def _normalise_image(data: bytes) -> bytes:
    with Image.open(io.BytesIO(data)) as opened:
        rgb = opened.convert("RGB")
        rgb.thumbnail((2000, 2000))
        return _png(rgb)


def _png(image: Image.Image) -> bytes:
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


def _tidy(text: str) -> str:
    text = text.replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()
